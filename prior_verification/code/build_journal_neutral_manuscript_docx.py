"""Build an editable BMC-style research manuscript from Markdown and verified figures."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


INK = RGBColor(0, 0, 0)
BLUE = RGBColor(0, 0, 0)
DARK_BLUE = RGBColor(0, 0, 0)
MUTED = RGBColor(70, 70, 70)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_font(run, size: float | None = None, color: RGBColor | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name, size, color, before, after in (
        ("Heading 1", 14, BLUE, 12, 6),
        ("Heading 2", 12, BLUE, 10, 4),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Caption" not in doc.styles:
        doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.bold = False
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0
    caption.paragraph_format.keep_with_next = False

    section_properties = section._sectPr
    line_numbers = section_properties.find(qn("w:lnNumType"))
    if line_numbers is None:
        line_numbers = OxmlElement("w:lnNumType")
        section_properties.append(line_numbers)
    line_numbers.set(qn("w:countBy"), "1")
    line_numbers.set(qn("w:start"), "1")
    line_numbers.set(qn("w:restart"), "continuous")
    line_numbers.set(qn("w:distance"), "360")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Audit before inference")
    set_font(header_run, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed = OxmlElement("w:t")
    displayed.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, displayed, end])


def plain_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = text.replace("--", "–")
    return text


def add_render_guard(doc: Document) -> None:
    """Submission format uses style spacing; do not add line-numbered spacer text."""
    return None


def add_body(doc: Document, draft: str) -> str:
    lines = draft.splitlines()
    title = ""
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("# "):
            title = plain_markdown(line[2:])
            index += 1
            break
        index += 1
    status = doc.add_paragraph()
    status.paragraph_format.space_after = Pt(6)
    status.paragraph_format.line_spacing = 1.0
    run = status.add_run("RESEARCH ARTICLE")
    set_font(run, size=10, color=MUTED, bold=True)
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(5)
    run = title_p.add_run(title)
    title_p.paragraph_format.line_spacing = 1.0
    set_font(run, size=18, color=INK, bold=True)

    for line in lines[index:]:
        if not line.strip():
            continue
        if line.startswith("## "):
            doc.add_paragraph(plain_markdown(line[3:]), style="Heading 1")
            add_render_guard(doc)
            continue
        if line.startswith("### "):
            doc.add_paragraph(plain_markdown(line[4:]), style="Heading 2")
            add_render_guard(doc)
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            run = paragraph.add_run(plain_markdown(line[2:]))
            set_font(run)
            add_render_guard(doc)
            continue
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(plain_markdown(line))
        set_font(run)
    return title


def load_figure(manifest_path: Path) -> dict:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    figure = Path(payload["outputs"]["png"]["path"])
    if payload["outputs"]["png"]["sha256"] != sha256(figure):
        raise ValueError(f"figure does not match its verified manifest: {manifest_path}")
    payload["_manifest_path"] = str(manifest_path)
    payload["_manifest_sha256"] = sha256(manifest_path)
    payload["_figure_path"] = str(figure)
    payload["_figure_sha256"] = sha256(figure)
    return payload


def caption_for(payload: dict) -> str:
    if "caption" in payload:
        return payload["caption"]
    source_hash = payload["source_tsv_sha256"]
    return (
        "Figure 1B. Executable gate outcomes. The figure is generated solely from the locked publication table "
        f"(SHA-256 {source_hash}). It displays technical gate status, row/LD counts, and numerical diagnostics only. "
        "READY indicates that supplied technical contracts passed; INELIGIBLE preserves a stop record. "
        "Neither status represents an association, posterior, biological result, or method-performance comparison."
    )


def add_figures(doc: Document, figures: list[dict]) -> None:
    for index, payload in enumerate(figures):
        if index == 0:
            doc.add_paragraph("Figures", style="Heading 1")
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        figure_width = 5.0 if index == 0 else 5.9
        picture = paragraph.add_run().add_picture(payload["_figure_path"], width=Inches(figure_width))
        alt_text = plain_markdown(payload.get("caption", f"Figure {index + 1}"))
        picture._inline.docPr.set("descr", alt_text[:1000])
        picture._inline.docPr.set("title", f"Figure {payload.get('figure_number', index + 1)}")
        caption = doc.add_paragraph(style="Caption")
        # Explicitly neutralise any inherited list indentation so every figure
        # legend begins within the page margin after the preceding callout list.
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # A small explicit inset prevents a LibreOffice edge-clipping defect
        # seen with long figure captions after a full-width inline SVG/PNG.
        caption.paragraph_format.left_indent = Inches(0.08)
        caption.paragraph_format.first_line_indent = Inches(0)
        caption.paragraph_format.right_indent = Inches(0.08)
        caption.paragraph_format.keep_together = True
        run = caption.add_run(caption_for(payload))
        set_font(run, size=10, color=MUTED, bold=False)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def set_table_three_line_base(table) -> None:
    """Remove grid styling so only explicit top/header/bottom rules remain."""
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is not None:
        properties.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    properties.append(borders)


def set_cell_rule(cell, edge: str, size: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    existing = borders.find(qn(f"w:{edge}"))
    if existing is not None:
        borders.remove(existing)
    node = OxmlElement(f"w:{edge}")
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), "000000")
    borders.append(node)


def apply_exact_table_geometry(table, widths_inches: list[float]) -> None:
    """Synchronize tblW, tblInd, tblGrid and every tcW in DXA."""
    widths = [int(round(float(width) * 1440)) for width in widths_inches]
    total = sum(widths)
    properties = table._tbl.tblPr
    table_width = ensure_child(properties, "w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    indent = ensure_child(properties, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "70")
    layout = ensure_child(properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            cell_width = ensure_child(cell._tc.get_or_add_tcPr(), "w:tcW")
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def load_tables(manifest_path: Path | None) -> list[dict]:
    if manifest_path is None:
        return []
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    tables = []
    for item in payload["tables"]:
        table_path = Path(item["path"])
        if not table_path.is_absolute():
            table_path = (manifest_path.parent / table_path).resolve()
        with table_path.open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle, delimiter="\t"))
        if len(rows) < 2 or len({len(row) for row in rows}) != 1:
            raise ValueError(f"table is empty or ragged: {table_path}")
        tables.append({**item, "_path": str(table_path), "_sha256": sha256(table_path), "_rows": rows})
    return tables


def add_tables(doc: Document, tables: list[dict]) -> None:
    if not tables:
        return
    doc.add_paragraph("Tables", style="Heading 1")
    for index, payload in enumerate(tables):
        if index:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(10)
            spacer.paragraph_format.space_after = Pt(0)
        heading = doc.add_paragraph()
        heading.paragraph_format.keep_with_next = True
        run = heading.add_run(f"Table {payload['number']}. {payload['title']}")
        set_font(run, size=11, color=INK, bold=True)
        table = doc.add_table(rows=len(payload["_rows"]), cols=len(payload["_rows"][0]))
        table.style = None
        widths = payload.get("column_widths_inches")
        if not widths:
            widths = [6.35 / len(payload["_rows"][0])] * len(payload["_rows"][0])
        apply_exact_table_geometry(table, [float(width) for width in widths])
        set_table_three_line_base(table)
        for r_index, source_row in enumerate(payload["_rows"]):
            for c_index, value in enumerate(source_row):
                cell = table.cell(r_index, c_index)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                run = paragraph.add_run(value)
                set_font(run, size=8 if len(source_row) >= 4 else 9, color=INK, bold=r_index == 0)
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
                run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
                if r_index == 0:
                    set_cell_rule(cell, "top", 12)
                    set_cell_rule(cell, "bottom", 8)
                if r_index == len(payload["_rows"]) - 1:
                    set_cell_rule(cell, "bottom", 12)
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_properties.append(repeat)
        caption = doc.add_paragraph(style="Caption")
        caption.paragraph_format.left_indent = Inches(0.08)
        caption.paragraph_format.right_indent = Inches(0.08)
        run = caption.add_run(payload["caption"])
        set_font(run, size=10, color=MUTED, bold=False)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown", required=True)
    parser.add_argument("--figure-manifests", required=True, nargs="+")
    parser.add_argument("--table-manifest")
    parser.add_argument("--out-docx", required=True)
    parser.add_argument("--manifest", required=True)
    args = parser.parse_args()
    markdown = Path(args.markdown).resolve()
    figure_manifests = [Path(item).resolve() for item in args.figure_manifests]
    out_docx = Path(args.out_docx).resolve()
    manifest = Path(args.manifest).resolve()
    figures = [load_figure(item) for item in figure_manifests]
    table_manifest = Path(args.table_manifest).resolve() if args.table_manifest else None
    tables = load_tables(table_manifest)
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    manifest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    title = add_body(doc, markdown.read_text(encoding="utf-8"))
    add_tables(doc, tables)
    add_figures(doc, figures)
    doc.core_properties.title = title
    doc.core_properties.author = "Author metadata required before submission"
    doc.core_properties.subject = "BMC Bioinformatics-style computational methods research article"
    doc.core_properties.comments = "Scientific content revised; author-owned declarations and public archive details remain required."
    doc.save(out_docx)
    payload = {
        "schema_version": "1.0",
        "document_style": "BMC research-manuscript override: Times New Roman 12 pt, double-spaced body, continuous line numbering, page numbering, black three-line tables",
        "submission_status": "SCIENTIFIC AND FORMAT REVISION COMPLETE — author metadata, declarations, licence, and public archive remain required",
        "markdown": {"path": str(markdown), "sha256": sha256(markdown)},
        "figures": [
            {
                "number": payload.get("figure_number", "1B"),
                "path": payload["_figure_path"],
                "sha256": payload["_figure_sha256"],
                "manifest": payload["_manifest_path"],
                "manifest_sha256": payload["_manifest_sha256"],
            }
            for payload in figures
        ],
        "tables": [
            {"number": payload["number"], "path": payload["_path"], "sha256": payload["_sha256"]}
            for payload in tables
        ],
        "docx": {"path": str(out_docx), "sha256": sha256(out_docx)},
        "author_actions_preserved": ["authors", "affiliations", "corresponding author", "institutional ethics wording", "funding", "competing interests", "author contributions", "code licence", "public repository/archive"],
    }
    manifest.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
