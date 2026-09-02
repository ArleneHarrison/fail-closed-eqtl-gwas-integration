"""Build a paginated DOCX review copy from the conservative revision sources."""
from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def set_run_font(run, size: float, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_document(path: Path) -> None:
    document = Document(path)
    for section in document.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.clear()
        set_run_font(header.add_run("Revision review copy"), 8, color="666666")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.clear()
        set_run_font(footer.add_run("Page "), 8, color="666666")
        add_page_field(footer)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in (("Title", 18, "17365D"), ("Heading 1", 14, "17365D"), ("Heading 2", 12, "2E5C88"), ("Heading 3", 10.5, "2E5C88")):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)

    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    for run in paragraph.runs:
                        set_run_font(run, 8.2, bold=row_index == 0)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "E8EEF5")
                    cell._tc.get_or_add_tcPr().append(shading)
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--pandoc", required=True, type=Path)
    parser.add_argument("--manuscript", required=True, type=Path)
    parser.add_argument("--captions", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    review_markdown = args.out.with_suffix(".md")
    figure_dir = args.manuscript.parent / "figures"
    review_markdown.write_text(args.manuscript.read_text(encoding="utf-8"), encoding="utf-8")
    args.out.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [str(args.pandoc), str(review_markdown), "--from", "markdown", "--to", "docx", "--output", str(args.out)],
        check=True,
        cwd=args.manuscript.parent,
    )
    style_document(args.out)


if __name__ == "__main__":
    main()
